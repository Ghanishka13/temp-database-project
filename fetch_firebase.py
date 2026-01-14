import firebase_admin
from firebase_admin import credentials, db
import pandas as pd
from openpyxl import load_workbook, Workbook
from datetime import datetime
import os
from dotenv import load_dotenv

load_dotenv()

EXCEL_FILE = "data.xlsx"
CSV_FILE = "data.csv"

def init_firebase():
    if not firebase_admin._apps:
        cred_dict = {
            "type": os.getenv("FIREBASE_TYPE"),
            "project_id": os.getenv("FIREBASE_PROJECT_ID"),
            "private_key_id": os.getenv("FIREBASE_PRIVATE_KEY_ID"),
            "private_key": os.getenv("FIREBASE_PRIVATE_KEY").replace('\\n', '\n'),
            "client_email": os.getenv("FIREBASE_CLIENT_EMAIL"),
            "client_id": os.getenv("FIREBASE_CLIENT_ID"),
            "auth_uri": os.getenv("FIREBASE_AUTH_URI"),
            "token_uri": os.getenv("FIREBASE_TOKEN_URI"),
            "auth_provider_x509_cert_url": os.getenv("FIREBASE_AUTH_PROVIDER_X509_CERT_URL"),
            "client_x509_cert_url": os.getenv("FIREBASE_CLIENT_X509_CERT_URL"),
        }
        cred = credentials.Certificate(cred_dict)
        firebase_admin.initialize_app(cred, {
            "databaseURL": os.getenv("FIREBASE_DATABASE_URL")
        })

def fetch_data():
    init_firebase()

    ref = db.reference("temperature")
    value = ref.get()

    if value is None:
        print("[ERROR] No data found.")
        return None

    # Extract actual temperature from nested data
    actual_temp = value.get("temperature")

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    return {
        "timestamp": timestamp,
        "temperature": actual_temp
    }

def save_to_csv(data):
    df = pd.DataFrame([data])
    df.to_csv(CSV_FILE, mode='a', header=not os.path.exists(CSV_FILE), index=False)
    print(f"[OK] Saved to CSV: {CSV_FILE}")

def save_to_excel(data):
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.append(["timestamp", "temperature"])
        wb.save(EXCEL_FILE)

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    ws.append([data["timestamp"], data["temperature"]])
    wb.save(EXCEL_FILE)
    print(f"[OK] Saved to Excel: {EXCEL_FILE}")

def main():
    data = fetch_data()
    if data:
        save_to_csv(data)
        save_to_excel(data)
        print("\nDATA SAVED SUCCESSFULLY\n")

if __name__ == "__main__":
    main()
